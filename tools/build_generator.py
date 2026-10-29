from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

MAX_MINES = 10
MAX_DAYS = 730
MAX_ROWS = MAX_MINES * MAX_DAYS
EVENTS = 6
OUT = Path(__file__).resolve().parents[1] / "generator" / "mining_generator.xlsx"
TITLE = Font(bold=True, size=16, color="FFFFFF")
HEADER = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
DARK = PatternFill("solid", fgColor="1F2A44")
MID = PatternFill("solid", fgColor="3B4A6B")
INPUT = PatternFill("solid", fgColor="FFF4CC")
LIGHT = PatternFill("solid", fgColor="EEF1F6")
THIN = Side(style="thin", color="BFC6D4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
DAYS_OF_WEEK = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
DOW_DEFAULT = [1, 1.02, 1.03, 1.01, 0.97, 0.75, 0.6]
MINES = ["LV-426", "Origae-6", "Fiorina 161", "Hadley's Hope", "Thedus"]
EVENT_ROWS = [
    (date(2026, 2, 10), 4, 0.35, 1.0, None),
    (date(2026, 3, 20), 2, 1.8, 0.5, None),
    (date(2026, 4, 25), 6, 0.6, None, 2),
    (date(2026, 5, 30), 3, 1.6, 0.3, None),
]
wb = Workbook()
def name(label, ref):
    wb.defined_names[label] = DefinedName(label, attr_text=ref)
def style(cell, font=None, fill=None, fmt=None, border=True, align=None):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = BOX
    if align:
        cell.alignment = align
    return cell
def put(ws, ref, value, **kw):
    ws[ref] = value
    return style(ws[ref], **kw)
def section(ws, row, text, width=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    put(ws, f"A{row}", text, font=HEADER, fill=MID)
def param(ws, row, label, value, fmt=None):
    put(ws, f"A{row}", label, font=BOLD, fill=LIGHT)
    put(ws, f"B{row}", value, fill=INPUT, fmt=fmt)
    return f"Generator!$B${row}"
def validation(ws, ref, **kw):
    dv = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(dv)
    dv.add(ref)
def build_generator():
    ws = wb.active
    ws.title = "Generator"
    ws.merge_cells("A1:H1")
    put(ws, "A1", "Weyland-Yutani Corporation  |  Mining Ops Simulator", font=TITLE, fill=DARK, align=CENTER)
    ws.row_dimensions[1].height = 30
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 15
    section(ws, 3, "Range and randomness")
    refs = {
        "Seed": param(ws, 4, "Seed", 42),
        "StartDate": param(ws, 5, "Start date", date(2026, 1, 1), fmt="yyyy-mm-dd"),
        "Days": param(ws, 6, f"Days (1-{MAX_DAYS})", 180),
        "Dist": param(ws, 7, "Distribution", "Normal"),
        "DistA": param(ws, 8, None, 1000),
        "DistB": param(ws, 9, None, 120),
        "Rho": param(ws, 10, "Smoothing / correlation (0-0.99)", 0.75),
        "Trend": param(ws, 11, "Trend (growth per day)", 0.001, fmt="0.00%"),
        "Spread": param(ws, 12, "Mine size spread", 0.3, fmt="0%"),
    }
    ws["A8"] = '=IF(Dist="Normal","Mean","Minimum")'
    ws["A9"] = '=IF(Dist="Normal","Standard deviation","Maximum")'
    validation(ws, "B6", type="whole", operator="between", formula1="1", formula2=str(MAX_DAYS))
    validation(ws, "B7", type="list", formula1='"Normal,Uniform"')
    validation(ws, "B10", type="decimal", operator="between", formula1="0", formula2="0.99")

    section(ws, 14, "Day-of-week factors")
    for i, (day, factor) in enumerate(zip(DAYS_OF_WEEK, DOW_DEFAULT)):
        col = get_column_letter(i + 2)
        put(ws, f"{col}15", day, font=BOLD, fill=LIGHT, align=CENTER)
        put(ws, f"{col}16", factor, fill=INPUT, fmt="0.00", align=CENTER)
    put(ws, "A16", "Output multiplier", font=BOLD, fill=LIGHT)
    refs["Dow"] = "Generator!$B$16:$H$16"
    section(ws, 18, f"Mines (fill from the top, up to {MAX_MINES})")
    for i in range(MAX_MINES):
        row = 19 + i
        put(ws, f"A{row}", f"Mine {i + 1}", font=BOLD, fill=LIGHT)
        put(ws, f"B{row}", MINES[i] if i < len(MINES) else None, fill=INPUT)
    refs["MineNames"] = f"Generator!$B$19:$B${18 + MAX_MINES}"
    count_row = 19 + MAX_MINES
    put(ws, f"A{count_row}", "Number of mines", font=BOLD, fill=LIGHT)
    put(ws, f"B{count_row}", "=COUNTA(MineNames)", font=BOLD)
    refs["MineCount"] = f"Generator!$B${count_row}"
    ev_head = count_row + 2
    section(ws, ev_head, "Events (spikes and drops, bell-shaped)")
    headers = ["Event", "Date", "Duration (days)", "Factor", "Probability", "Mines affected"]
    for i, text in enumerate(headers):
        put(ws, f"{get_column_letter(i + 1)}{ev_head + 1}", text, font=BOLD, fill=LIGHT, align=CENTER)
    first = ev_head + 2
    for i in range(EVENTS):
        row = first + i
        values = EVENT_ROWS[i] if i < len(EVENT_ROWS) else (None,) * 5
        put(ws, f"A{row}", f"Event {i + 1}", font=BOLD, fill=LIGHT)
        for col, value, fmt in zip("BCDEF", values, ["yyyy-mm-dd", "0", "0.00", "0%", "0"]):
            put(ws, f"{col}{row}", value, fill=INPUT, fmt=fmt, align=CENTER)
    last = first + EVENTS - 1
    for label, col in zip(["EvDate", "EvDur", "EvFactor", "EvProb", "EvN"], "BCDEF"):
        refs[label] = f"Generator!${col}${first}:${col}${last}"
    validation(ws, f"E{first}:E{last}", type="decimal", operator="between", formula1="0", formula2="1")
    validation(ws, f"F{first}:F{last}", type="whole", operator="between", formula1="0", formula2=str(MAX_MINES))
    note = last + 2
    ws.merge_cells(start_row=note, start_column=1, end_row=note + 2, end_column=8)
    put(ws, f"A{note}", (
        "Yellow cells are inputs. Factor < 1 is a drop, > 1 is a spike. "
        "Probability is the chance that a mine is affected; if 'Mines affected' is set, "
        "exactly that many randomly chosen mines are affected instead. "
        "Results are reproducible for the same seed."
    ), border=False, align=Alignment(wrap_text=True, vertical="top"))
    for label, ref in refs.items():
        name(label, ref)
def build_helpers():
    ws = wb.create_sheet("Helpers")
    put(ws, "A1", "Mine", font=HEADER, fill=MID)
    put(ws, "B1", "Name", font=HEADER, fill=MID)
    put(ws, "C1", "Size factor", font=HEADER, fill=MID)
    for m in range(1, MAX_MINES + 1):
        r = m + 1
        ws[f"A{r}"] = m
        ws[f"B{r}"] = f'=IF(A{r}<=MineCount,INDEX(MineNames,A{r}),"")'
        ws[f"C{r}"] = f"=1+Spread*(2*MOD(SIN(Seed*1.3711+A{r}*91.733)*43758.5453,1)-1)"
    name("MineScale", f"Helpers!$C$2:$C${MAX_MINES + 1}")
    top = MAX_MINES + 3
    rand_cols = [get_column_letter(5 + m) for m in range(MAX_MINES)]
    aff_cols = [get_column_letter(5 + MAX_MINES + m) for m in range(MAX_MINES)]
    for col, text in zip("ABCD", ["Event", "Factor", "Center", "Sigma"]):
        put(ws, f"{col}{top}", text, font=HEADER, fill=MID)
    for m, (rc, ac) in enumerate(zip(rand_cols, aff_cols), 1):
        put(ws, f"{rc}{top}", m, font=HEADER, fill=MID)
        put(ws, f"{ac}{top}", m, font=HEADER, fill=MID)
    put(ws, f"{rand_cols[0]}{top - 1}", "Random rank key per mine", font=BOLD, border=False)
    put(ws, f"{aff_cols[0]}{top - 1}", "Affected (1) or not (0)", font=BOLD, border=False)
    for e in range(1, EVENTS + 1):
        r = top + e
        ev = f"INDEX({{}},{e})"
        ws[f"A{r}"] = e
        ws[f"B{r}"] = f'=IF({ev.format("EvDate")}="",1,MAX({ev.format("EvFactor")},0.01))'
        ws[f"C{r}"] = f'=IF({ev.format("EvDate")}="",0,{ev.format("EvDate")}+(MAX({ev.format("EvDur")},1)-1)/2)'
        ws[f"D{r}"] = f'=MAX({ev.format("EvDur")},1)/2'
        rand_row = f"${rand_cols[0]}{r}:${rand_cols[-1]}{r}"
        for m, (rc, ac) in enumerate(zip(rand_cols, aff_cols), 1):
            ws[f"{rc}{r}"] = f"=IF({m}>MineCount,2,MOD(SIN(Seed*0.9137+{e}*311.71+{m}*17.317)*43758.5453,1))"
            ws[f"{ac}{r}"] = (
                f'=IF(OR({m}>MineCount,{ev.format("EvDate")}=""),0,'
                f'IF({ev.format("EvN")}<>"",IF(COUNTIF({rand_row},"<"&{rc}{r})<{ev.format("EvN")},1,0),'
                f'IF({rc}{r}<{ev.format("EvProb")},1,0)))'
            )
    name("EvEff", f"Helpers!$B${top + 1}:$B${top + EVENTS}")
    name("EvCenter", f"Helpers!$C${top + 1}:$C${top + EVENTS}")
    name("EvSigma", f"Helpers!$D${top + 1}:$D${top + EVENTS}")
    name("Affected", f"Helpers!${aff_cols[0]}${top + 1}:${aff_cols[-1]}${top + EVENTS}")
CALC_COLUMNS = [
    ("k", None),
    ("active", "=A{r}<=MineCount*Days"),
    ("mine", '=IF(B{r},INT((A{r}-1)/Days)+1,"")'),
    ("day", '=IF(B{r},MOD(A{r}-1,Days),"")'),
    ("date", '=IF(B{r},StartDate+D{r},"")'),
    ("u1", '=IF(B{r},MAX(MOD(SIN(Seed*0.1373+C{r}*12.9898+D{r}*78.233)*43758.5453,1),1E-9),"")'),
    ("u2", '=IF(B{r},MOD(SIN(Seed*0.5917+C{r}*39.3468+D{r}*11.135)*24634.6345,1),"")'),
    ("z", '=IF(B{r},SQRT(-2*LN(F{r}))*COS(2*PI()*G{r}),"")'),
    ("smoothed", '=IF(B{r},IF(D{r}=0,H{r},Rho*I{p}+SQRT(1-Rho^2)*H{r}),"")'),
    ("base", '=IF(B{r},IF(Dist="Normal",DistA+DistB*I{r},DistA+(DistB-DistA)*NORMSDIST(I{r})),"")'),
    ("size", '=IF(B{r},INDEX(MineScale,C{r}),"")'),
    ("weekday", '=IF(B{r},INDEX(Dow,1,WEEKDAY(E{r},2)),"")'),
    ("trend", '=IF(B{r},(1+Trend)^D{r},"")'),
    ("events", '=IF(B{r},EXP(SUMPRODUCT(LN(1+(EvEff-1)*INDEX(Affected,0,C{r})*EXP(-((E{r}-EvCenter)^2)/(2*EvSigma^2))))),"")'),
    ("output", '=IF(B{r},ROUND(MAX(0,J{r}*K{r}*L{r}*M{r}*N{r}),2),"")'),
]
def build_calc():
    ws = wb.create_sheet("Calc")
    for c, (title, _) in enumerate(CALC_COLUMNS, 1):
        put(ws, f"{get_column_letter(c)}1", title, font=HEADER, fill=MID)
    for k in range(1, MAX_ROWS + 1):
        r = k + 1
        ws.cell(row=r, column=1, value=k)
        for c, (_, formula) in enumerate(CALC_COLUMNS[1:], 2):
            ws.cell(row=r, column=c, value=formula.format(r=r, p=r - 1))
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
def build_data():
    ws = wb.create_sheet("Data", 1)
    for col, text, width in [("A", "Date", 14), ("B", "Mine", 18), ("C", "Output", 12)]:
        put(ws, f"{col}1", text, font=HEADER, fill=DARK)
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    for k in range(1, MAX_ROWS + 1):
        r = k + 1
        ws[f"A{r}"] = f'=IF(Calc!$B{r},Calc!E{r},"")'
        ws[f"B{r}"] = f'=IF(Calc!$B{r},INDEX(MineNames,Calc!C{r}),"")'
        ws[f"C{r}"] = f'=IF(Calc!$B{r},Calc!O{r},"")'
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
        ws[f"C{r}"].number_format = "0.00"
def build_chart():
    wb.create_sheet("Chart", 2)
build_generator()
build_helpers()
build_calc()
build_data()
build_chart()
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)